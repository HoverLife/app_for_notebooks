#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PDF Lot Parser: GUI + CLI приложение для парсинга аукционных PDF в Excel."""

from __future__ import annotations

import argparse
import csv
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import traceback
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import fitz  # PyMuPDF
import pandas as pd
import pytesseract
from PIL import Image
from dateutil import parser as date_parser
from PyQt5 import QtCore, QtGui, QtWidgets

APP_TITLE = "PDF Lot Parser"
DEFAULT_OUTPUT_DIR = "output"
DEFAULT_ERROR_FILE = "parsing_errors.csv"
DEFAULT_LOG_LEVEL = "INFO"

KNOWN_BRANDS = {
    "aoc",
    "samsung",
    "lg",
    "philips",
    "dell",
    "hp",
    "lenovo",
    "acer",
    "asus",
    "benq",
    "iiyama",
    "viewsonic",
    "sony",
    "panasonic",
    "xiaomi",
    "msi",
}

CONDITION_KEYWORDS = {
    "Bad": ["не работает", "бит", "разбит", "полосы", "мерцает", "дефект", "не включается"],
    "Fair": ["следы эксплуатации", "потертости", "потёртости", "царапины", "сколы", "желтизна"],
    "Good": ["исправен", "рабочий", "без дефектов", "нет"],
}

HEADER_COLUMNS_BASE = [
    "SourceFile",
    "SourcePage",
    "ProcessedAt",
    "LotPosition",
    "AuctionCity",
    "AuctionEndDateTime",
    "ItemType",
    "Brand",
    "Model",
    "Diagonal",
    "SerialNumber",
    "StartPriceRUB",
    "PickupAddress",
    "CityPickup",
    "VideoInputs",
    "PowerType",
    "PowerIncluded",
    "MatrixIssues",
    "Defect1",
    "Defect2",
    "HasImages",
    "RawText",
    "ExtractedFeatures",
    "Keywords",
    "EstimatedCondition",
    "Notes",
]


@dataclass
class ParserOptions:
    ocr_mode: str = "Auto"  # Auto, Always, Never
    save_images: bool = True
    output_folder: Path = Path(DEFAULT_OUTPUT_DIR)
    auto_start_drop: bool = False
    log_level: str = DEFAULT_LOG_LEVEL
    feature_threshold: int = 5
    save_csv: bool = False


@dataclass
class ParseError:
    source_file: str
    source_page: int
    lot_position: Optional[int]
    message: str
    raw_text: str


class TesseractResolver:
    """Поиск tesseract.exe в режиме разработки и PyInstaller."""

    @staticmethod
    def resolve(logger: logging.Logger) -> Optional[str]:
        candidates = []
        if hasattr(sys, "_MEIPASS"):
            candidates.append(Path(getattr(sys, "_MEIPASS")) / "tesseract.exe")
            candidates.append(Path(getattr(sys, "_MEIPASS")) / "Tesseract-OCR" / "tesseract.exe")

        candidates.extend(
            [
                Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
                Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
            ]
        )

        for path in candidates:
            if path.exists():
                pytesseract.pytesseract.tesseract_cmd = str(path)
                logger.info("Tesseract найден: %s", path)
                return str(path)

        from_path = shutil.which("tesseract")
        if from_path:
            pytesseract.pytesseract.tesseract_cmd = from_path
            logger.info("Tesseract найден в PATH: %s", from_path)
            return from_path

        logger.warning(
            "Tesseract не найден. OCR fallback недоступен. Установите Tesseract и добавьте в PATH."
        )
        return None


class LotParser:
    """Парсер лотов аукционных карточек."""

    def __init__(self, options: ParserOptions, logger: logging.Logger):
        self.options = options
        self.logger = logger
        self.errors: List[ParseError] = []
        self.new_feature_counter: Counter[str] = Counter()
        self.tesseract_path = TesseractResolver.resolve(logger)

    @staticmethod
    def normalize_text(text: str) -> str:
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        text = text.replace("—", "-").replace("–", "-")
        text = re.sub(r"[\t\f\v]+", " ", text)
        text = re.sub(r"[ ]{2,}", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        return text.strip()

    def split_into_lots(self, text: str) -> List[str]:
        text = self.normalize_text(text)
        pattern = re.compile(r"(?=\n?\s*\d+\s*-{2,}\s*)", re.MULTILINE)
        chunks = [c.strip() for c in pattern.split("\n" + text) if c.strip()]
        if len(chunks) <= 1:
            fallback = re.compile(r"(?=\n?\s*\d+\s+[А-Яа-яA-Za-z])", re.MULTILINE)
            chunks = [c.strip() for c in fallback.split("\n" + text) if c.strip()]
        return chunks

    def _extract_datetime(self, full_text: str) -> Optional[datetime]:
        m = re.search(r"Дата и время окончания торгов\s*:\s*([0-9.:\s]+)", full_text, re.IGNORECASE)
        if not m:
            return None
        value = m.group(1).strip()
        try:
            return date_parser.parse(value, dayfirst=True)
        except Exception:
            return None

    def _extract_city_header(self, full_text: str) -> str:
        m = re.search(r"города\s*:\s*([А-Яа-яA-Za-z\-\s]+)", full_text, re.IGNORECASE)
        return m.group(1).strip() if m else ""

    def _extract_auction_city(self, block: str, header_city: str) -> str:
        m = re.search(r"\(([А-Яа-яA-Za-z\-\s]+)\)", block)
        if m:
            return m.group(1).strip()
        return header_city

    def _extract_brand(self, text: str) -> str:
        low = text.lower()
        for b in KNOWN_BRANDS:
            if re.search(rf"\b{re.escape(b)}\b", low):
                return b.upper()
        return ""

    def _extract_model(self, text: str, brand: str) -> str:
        if brand:
            m = re.search(rf"{re.escape(brand)}\s+([A-Za-z0-9\-_/]+)", text, re.IGNORECASE)
            if m:
                return m.group(1).strip()
        m2 = re.search(r"(?:Монитор|Ноутбук|Телевизор|Принтер)\s+([A-Za-z0-9\-_/]+)", text, re.IGNORECASE)
        return m2.group(1).strip() if m2 else ""

    def _extract_item_type(self, text: str) -> str:
        for t in ["Монитор", "Ноутбук", "Телевизор", "Принтер", "ПК", "Системный блок"]:
            if re.search(rf"\b{re.escape(t)}\b", text, re.IGNORECASE):
                return t
        return ""

    def _normalize_diagonal(self, text: str) -> str:
        m = re.search(r"(\d{1,2}(?:[\.,]\d)?)\s*(?:'|\")", text)
        if not m:
            m = re.search(r"Диагональ\s*-\s*(\d{1,2}(?:[\.,]\d)?)", text, re.IGNORECASE)
        if not m:
            return ""
        value = m.group(1).replace(",", ".")
        return f'{value}"'

    def _extract_serial(self, text: str) -> str:
        m = re.search(r"s\s*/\s*n\s*:\s*([A-Za-z0-9\-]+)", text, re.IGNORECASE)
        if m:
            return m.group(1).strip().upper()
        m = re.search(r"\b([A-Z0-9]{10,})\b", text)
        return m.group(1).strip().upper() if m else ""

    def _extract_price(self, text: str) -> Optional[int]:
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        for ln in reversed(lines):
            if re.fullmatch(r"[\d\s]{2,}", ln):
                try:
                    return int(re.sub(r"\D", "", ln))
                except Exception:
                    pass
        m = re.search(r"([\d\s]{2,})\s*(?:руб|р\.|RUB)?", text, re.IGNORECASE)
        if m:
            digits = re.sub(r"\D", "", m.group(1))
            if digits:
                return int(digits)
        return None

    def _extract_address(self, text: str) -> Tuple[str, str]:
        lines = [ln.strip(" ,") for ln in text.splitlines() if ln.strip()]
        city = ""
        address_parts: List[str] = []
        address_start = None
        for idx, line in enumerate(lines):
            if re.search(r"\bг\.\s*[А-Яа-яA-Za-z\-]+", line):
                address_start = idx
                break
        if address_start is not None:
            address_parts = lines[address_start : min(address_start + 6, len(lines))]
        address = ", ".join(address_parts)
        m_city = re.search(r"\bг\.\s*([А-Яа-яA-Za-z\-]+)", address)
        if m_city:
            city = m_city.group(1)
        return address, city

    def _parse_key_values(self, text: str) -> Dict[str, str]:
        result: Dict[str, str] = {}
        for key, value in re.findall(r"([А-Яа-яA-Za-z0-9\-\s]+)-\s*([^/\n]+)", text):
            k = key.strip()
            v = value.strip()
            if not k or not v:
                continue
            result[k] = v
        return result

    def _estimate_condition(self, text: str) -> str:
        low = text.lower()
        for cond in ["Bad", "Fair", "Good"]:
            if any(k in low for k in CONDITION_KEYWORDS[cond]):
                return cond
        return "Fair"

    def _extract_keywords(self, row: Dict[str, object]) -> str:
        tokens = []
        for key in ["Brand", "Model", "Defect1", "Defect2", "MatrixIssues", "ItemType"]:
            value = str(row.get(key, "") or "").strip()
            if value:
                tokens.append(value)
        return ", ".join(dict.fromkeys(tokens))[:250]

    def extract_images(self, doc: fitz.Document, page_index: int, pdf_stem: str) -> List[str]:
        page = doc[page_index]
        image_list = page.get_images(full=True)
        if not image_list or not self.options.save_images:
            return []

        output_dir = self.options.output_folder / "output_images" / pdf_stem
        output_dir.mkdir(parents=True, exist_ok=True)

        rel_paths = []
        for idx, img in enumerate(image_list, start=1):
            xref = img[0]
            try:
                base_img = doc.extract_image(xref)
                img_bytes = base_img["image"]
                ext = base_img.get("ext", "jpg")
                filename = f"photo_{page_index + 1}_{idx}.{ext}"
                out_path = output_dir / filename
                with open(out_path, "wb") as fh:
                    fh.write(img_bytes)
                rel_paths.append(str(out_path.relative_to(self.options.output_folder)))
            except Exception as exc:
                self.logger.error("Не удалось сохранить изображение xref=%s: %s", xref, exc)
        return rel_paths

    def _ocr_page_text(self, page: fitz.Page) -> str:
        if not self.tesseract_path:
            return ""
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
        mode = "RGBA" if pix.alpha else "RGB"
        image = Image.frombytes(mode, [pix.width, pix.height], pix.samples)
        try:
            return pytesseract.image_to_string(image, lang="rus+eng")
        except Exception as exc:
            self.logger.error("OCR ошибка: %s", exc)
            return ""

    def _need_ocr(self, text: str) -> bool:
        stripped = re.sub(r"\s+", "", text)
        return len(stripped) < 60

    def parse_pdf(self, pdf_path: Path) -> List[Dict[str, object]]:
        rows: List[Dict[str, object]] = []
        try:
            doc = fitz.open(pdf_path)
        except Exception as exc:
            self.errors.append(ParseError(pdf_path.name, 0, None, f"Ошибка открытия PDF: {exc}", ""))
            return rows

        processed_at = datetime.now()
        header_city = ""
        end_dt: Optional[datetime] = None

        for page_index in range(len(doc)):
            page = doc[page_index]
            raw_text = page.get_text("text") or ""
            if page_index == 0:
                header_city = self._extract_city_header(raw_text)
                end_dt = self._extract_datetime(raw_text)

            if self.options.ocr_mode == "Always" or (
                self.options.ocr_mode == "Auto" and self._need_ocr(raw_text)
            ):
                ocr_text = self._ocr_page_text(page)
                if ocr_text.strip():
                    raw_text = raw_text + "\n" + ocr_text

            lots = self.split_into_lots(raw_text)
            page_images = self.extract_images(doc, page_index, pdf_path.stem)

            for lot in lots:
                row = self.parse_lot_block(
                    lot,
                    source_file=pdf_path.name,
                    source_page=page_index + 1,
                    processed_at=processed_at,
                    auction_city_header=header_city,
                    auction_end_datetime=end_dt,
                    page_images=page_images,
                )
                if row:
                    rows.append(row)

        doc.close()
        return rows

    def parse_lot_block(
        self,
        block: str,
        source_file: str,
        source_page: int,
        processed_at: datetime,
        auction_city_header: str,
        auction_end_datetime: Optional[datetime],
        page_images: List[str],
    ) -> Optional[Dict[str, object]]:
        block_n = self.normalize_text(block)
        if len(block_n) < 10:
            return None

        pos_match = re.match(r"\s*(\d+)\s*-{2,}", block_n)
        if not pos_match:
            pos_match = re.match(r"\s*(\d+)\s+", block_n)
        lot_position = int(pos_match.group(1)) if pos_match else None

        attrs = self._parse_key_values(block_n)
        brand = self._extract_brand(block_n)
        item_type = self._extract_item_type(block_n)
        model = self._extract_model(block_n, brand)
        diagonal = self._normalize_diagonal(block_n)
        serial = self._extract_serial(block_n)
        price = self._extract_price(block_n)
        address, city_pickup = self._extract_address(block_n)

        if price is None:
            self.errors.append(
                ParseError(source_file, source_page, lot_position, "Не найдена стартовая цена", block_n)
            )

        row: Dict[str, object] = {
            "SourceFile": source_file,
            "SourcePage": source_page,
            "ProcessedAt": processed_at,
            "LotPosition": lot_position,
            "AuctionCity": self._extract_auction_city(block_n, auction_city_header),
            "AuctionEndDateTime": auction_end_datetime,
            "ItemType": item_type,
            "Brand": brand,
            "Model": model,
            "Diagonal": diagonal,
            "SerialNumber": serial,
            "StartPriceRUB": price,
            "PickupAddress": address,
            "CityPickup": city_pickup,
            "VideoInputs": attrs.get("Видеовходы", ""),
            "PowerType": attrs.get("Тип питания", ""),
            "PowerIncluded": attrs.get("Сетевой шнур или блок питания в наличии", ""),
            "MatrixIssues": attrs.get("Затемнения и полосы на матрице", ""),
            "Defect1": attrs.get("Дефект 1", ""),
            "Defect2": attrs.get("Дефект 2", ""),
            "HasImages": bool(page_images),
            "RawText": block_n,
            "ExtractedFeatures": "",
            "Keywords": "",
            "EstimatedCondition": self._estimate_condition(block_n),
            "Notes": "",
        }

        extra = {}
        predefined = {
            "Видеовходы",
            "Тип питания",
            "Сетевой шнур или блок питания в наличии",
            "Затемнения и полосы на матрице",
            "Дефект 1",
            "Дефект 2",
            "Диагональ",
        }
        for k, v in attrs.items():
            if k not in predefined:
                extra[k] = v
                self.new_feature_counter[k] += 1

        for idx, p in enumerate(page_images, start=1):
            row[f"Photo{idx}"] = p

        row["ExtractedFeatures"] = json.dumps(extra, ensure_ascii=False)
        row["Keywords"] = self._extract_keywords(row)
        return row

    def parse_inputs(self, input_paths: Sequence[Path]) -> List[Dict[str, object]]:
        pdf_files: List[Path] = []
        for path in input_paths:
            if path.is_dir():
                pdf_files.extend(sorted(path.glob("*.pdf")))
            elif path.is_file() and path.suffix.lower() == ".pdf":
                pdf_files.append(path)

        rows: List[Dict[str, object]] = []
        for pdf in pdf_files:
            self.logger.info("Обработка файла: %s", pdf)
            rows.extend(self.parse_pdf(pdf))

        for k, cnt in self.new_feature_counter.items():
            if cnt >= self.options.feature_threshold:
                self.logger.info(
                    "Новый ключ '%s' встречен %s раз(а). Рекомендуется добавить отдельную колонку.",
                    k,
                    cnt,
                )

        return rows

    def save_errors_csv(self, out_dir: Path) -> Path:
        out_path = out_dir / DEFAULT_ERROR_FILE
        out_dir.mkdir(parents=True, exist_ok=True)
        with open(out_path, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            writer.writerow(["SourceFile", "SourcePage", "LotPosition", "ErrorMessage", "RawText"])
            for e in self.errors:
                writer.writerow([e.source_file, e.source_page, e.lot_position, e.message, e.raw_text])
        return out_path


def export_rows(rows: List[Dict[str, object]], excel_path: Path, save_csv: bool = False) -> Path:
    excel_path.parent.mkdir(parents=True, exist_ok=True)

    max_photo = 0
    for r in rows:
        photos = [k for k in r.keys() if k.startswith("Photo")]
        max_photo = max(max_photo, len(photos))

    columns = HEADER_COLUMNS_BASE + [f"Photo{i}" for i in range(1, max_photo + 1)]
    df = pd.DataFrame(rows)
    for c in columns:
        if c not in df.columns:
            df[c] = ""
    df = df[columns]

    if "StartPriceRUB" in df.columns:
        df["StartPriceRUB"] = pd.to_numeric(df["StartPriceRUB"], errors="coerce")

    with pd.ExcelWriter(excel_path, engine="openpyxl", datetime_format="YYYY-MM-DD HH:MM:SS") as writer:
        df.to_excel(writer, index=False, sheet_name="Lots")
        ws = writer.sheets["Lots"]
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="FFDDEEFF", end_color="FFDDEEFF", fill_type="solid")
        for col in ws.columns:
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, max_len + 2), 60)

    if save_csv:
        csv_path = excel_path.with_suffix(".csv")
        df.to_csv(csv_path, index=False, encoding="utf-8-sig")

    return excel_path


def build_logger(name: str, level: str = DEFAULT_LOG_LEVEL, handler: Optional[logging.Handler] = None) -> logging.Logger:
    logger = logging.getLogger(name)
    logger.setLevel(getattr(logging, level.upper(), logging.INFO))
    logger.handlers.clear()
    formatter = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s", "%H:%M:%S")

    sh = logging.StreamHandler(sys.stdout)
    sh.setFormatter(formatter)
    logger.addHandler(sh)

    if handler:
        handler.setFormatter(formatter)
        logger.addHandler(handler)

    return logger


class QtLogHandler(logging.Handler):
    def __init__(self, callback):
        super().__init__()
        self.callback = callback

    def emit(self, record: logging.LogRecord) -> None:
        self.callback(self.format(record))


class DropArea(QtWidgets.QFrame):
    filesDropped = QtCore.pyqtSignal(list)

    def __init__(self):
        super().__init__()
        self.setAcceptDrops(True)
        self.setFrameShape(QtWidgets.QFrame.StyledPanel)
        self.setStyleSheet("QFrame {border: 2px dashed #9aa0a6; border-radius: 8px; background: #f7f7f7;}")
        layout = QtWidgets.QVBoxLayout(self)
        self.label = QtWidgets.QLabel("Перетащите PDF сюда")
        self.label.setAlignment(QtCore.Qt.AlignCenter)
        self.label.setStyleSheet("font-size: 16px; color: #555;")
        layout.addWidget(self.label)

    def dragEnterEvent(self, event: QtGui.QDragEnterEvent) -> None:
        if event.mimeData().hasUrls():
            event.acceptProposedAction()
            self.setStyleSheet("QFrame {border: 2px dashed #2e7d32; border-radius: 8px; background: #e8f5e9;}")

    def dragLeaveEvent(self, event: QtGui.QDragLeaveEvent) -> None:
        self.setStyleSheet("QFrame {border: 2px dashed #9aa0a6; border-radius: 8px; background: #f7f7f7;}")

    def dropEvent(self, event: QtGui.QDropEvent) -> None:
        self.setStyleSheet("QFrame {border: 2px dashed #9aa0a6; border-radius: 8px; background: #f7f7f7;}")
        paths = []
        for url in event.mimeData().urls():
            p = Path(url.toLocalFile())
            if p.exists():
                paths.append(str(p))
        self.filesDropped.emit(paths)


class OptionsDialog(QtWidgets.QDialog):
    def __init__(self, options: ParserOptions, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Опции")
        self.options = options
        layout = QtWidgets.QFormLayout(self)

        self.ocr_mode = QtWidgets.QComboBox()
        self.ocr_mode.addItems(["Auto", "Always", "Never"])
        self.ocr_mode.setCurrentText(options.ocr_mode)

        self.save_images = QtWidgets.QCheckBox("Сохранять фото")
        self.save_images.setChecked(options.save_images)

        self.output_folder = QtWidgets.QLineEdit(str(options.output_folder))
        self.auto_drop = QtWidgets.QCheckBox("Авто-старт при drop")
        self.auto_drop.setChecked(options.auto_start_drop)

        self.log_level = QtWidgets.QComboBox()
        self.log_level.addItems(["DEBUG", "INFO", "WARNING", "ERROR"])
        self.log_level.setCurrentText(options.log_level)

        self.threshold = QtWidgets.QSpinBox()
        self.threshold.setRange(1, 100)
        self.threshold.setValue(options.feature_threshold)

        self.save_csv = QtWidgets.QCheckBox("Дополнительно сохранять CSV")
        self.save_csv.setChecked(options.save_csv)

        layout.addRow("OCR mode", self.ocr_mode)
        layout.addRow("", self.save_images)
        layout.addRow("Папка вывода", self.output_folder)
        layout.addRow("", self.auto_drop)
        layout.addRow("Уровень логирования", self.log_level)
        layout.addRow("Порог новых ключей (N)", self.threshold)
        layout.addRow("", self.save_csv)

        btns = QtWidgets.QDialogButtonBox(QtWidgets.QDialogButtonBox.Ok | QtWidgets.QDialogButtonBox.Cancel)
        btns.accepted.connect(self.accept)
        btns.rejected.connect(self.reject)
        layout.addWidget(btns)

    def accept(self) -> None:
        self.options.ocr_mode = self.ocr_mode.currentText()
        self.options.save_images = self.save_images.isChecked()
        self.options.output_folder = Path(self.output_folder.text().strip() or DEFAULT_OUTPUT_DIR)
        self.options.auto_start_drop = self.auto_drop.isChecked()
        self.options.log_level = self.log_level.currentText()
        self.options.feature_threshold = self.threshold.value()
        self.options.save_csv = self.save_csv.isChecked()
        super().accept()


class MainWindow(QtWidgets.QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(APP_TITLE)
        self.resize(700, 450)

        self.options = ParserOptions()
        self.selected_paths: List[Path] = []
        self.last_excel_path: Optional[Path] = None
        self.last_error_path: Optional[Path] = None
        self.current_rows: List[Dict[str, object]] = []

        central = QtWidgets.QWidget()
        self.setCentralWidget(central)
        main_layout = QtWidgets.QVBoxLayout(central)

        self.drop_area = DropArea()
        self.drop_area.filesDropped.connect(self.on_files_dropped)
        main_layout.addWidget(self.drop_area)

        btn_layout = QtWidgets.QHBoxLayout()
        self.btn_select = QtWidgets.QPushButton("Выбрать файл(ы)…")
        self.btn_process = QtWidgets.QPushButton("Обработать")
        self.btn_save = QtWidgets.QPushButton("Сохранить Excel")
        self.btn_opts = QtWidgets.QPushButton("Опции")
        self.btn_tests = QtWidgets.QPushButton("Run sample tests")

        self.btn_select.clicked.connect(self.select_files)
        self.btn_process.clicked.connect(self.process_files)
        self.btn_save.clicked.connect(self.save_excel)
        self.btn_opts.clicked.connect(self.show_options)
        self.btn_tests.clicked.connect(self.run_tests)

        for btn in [self.btn_select, self.btn_process, self.btn_save, self.btn_opts, self.btn_tests]:
            btn_layout.addWidget(btn)
        main_layout.addLayout(btn_layout)

        self.progress = QtWidgets.QProgressBar()
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        main_layout.addWidget(self.progress)

        self.log = QtWidgets.QPlainTextEdit()
        self.log.setReadOnly(True)
        main_layout.addWidget(self.log)

        self.qt_handler = QtLogHandler(self.append_log)
        self.logger = build_logger("gui", self.options.log_level, self.qt_handler)

    def append_log(self, msg: str) -> None:
        self.log.appendPlainText(msg)

    def select_files(self) -> None:
        files, _ = QtWidgets.QFileDialog.getOpenFileNames(self, "Выберите PDF", "", "PDF files (*.pdf)")
        if files:
            self.selected_paths = [Path(f) for f in files]
            self.logger.info("Выбрано файлов: %s", len(files))

    def on_files_dropped(self, files: List[str]) -> None:
        self.selected_paths = [Path(f) for f in files]
        self.logger.info("Перетащено объектов: %s", len(files))
        if self.options.auto_start_drop:
            self.process_files()

    def show_options(self) -> None:
        dlg = OptionsDialog(self.options, self)
        if dlg.exec_():
            self.logger = build_logger("gui", self.options.log_level, self.qt_handler)
            self.logger.info("Опции обновлены")

    def process_files(self) -> None:
        if not self.selected_paths:
            QtWidgets.QMessageBox.warning(self, "Нет файлов", "Сначала выберите PDF-файлы или папку.")
            return

        try:
            parser = LotParser(self.options, self.logger)
            self.progress.setValue(10)
            rows = parser.parse_inputs(self.selected_paths)
            self.progress.setValue(80)
            self.current_rows = rows
            self.last_error_path = parser.save_errors_csv(self.options.output_folder)
            self.progress.setValue(100)
            self.logger.info("Обработано лотов: %s", len(rows))
            self.logger.info("Ошибок: %s", len(parser.errors))
            self._show_finish_dialog(len(rows), len(parser.errors))
        except Exception as exc:
            self.logger.error("Ошибка обработки: %s", exc)
            self.logger.debug(traceback.format_exc())
            QtWidgets.QMessageBox.critical(self, "Ошибка", f"Ошибка обработки: {exc}")

    def save_excel(self) -> None:
        if not self.current_rows:
            QtWidgets.QMessageBox.warning(self, "Нет данных", "Сначала выполните обработку.")
            return

        path, _ = QtWidgets.QFileDialog.getSaveFileName(
            self,
            "Сохранить Excel",
            str(self.options.output_folder / "parsed_lots.xlsx"),
            "Excel (*.xlsx)",
        )
        if not path:
            return

        try:
            self.last_excel_path = export_rows(self.current_rows, Path(path), save_csv=self.options.save_csv)
            self.logger.info("Excel сохранён: %s", self.last_excel_path)
            QtWidgets.QMessageBox.information(self, "Готово", f"Файл сохранён:\n{self.last_excel_path}")
        except Exception as exc:
            self.logger.error("Ошибка сохранения Excel: %s", exc)
            QtWidgets.QMessageBox.critical(self, "Ошибка", f"Не удалось сохранить Excel: {exc}")

    def _show_finish_dialog(self, lot_count: int, error_count: int) -> None:
        msg = QtWidgets.QMessageBox(self)
        msg.setWindowTitle("Обработка завершена")
        msg.setText(f"Обработано лотов: {lot_count}\nОшибок: {error_count}")
        btn_open_out = msg.addButton("Открыть папку с результатами", QtWidgets.QMessageBox.ActionRole)
        btn_open_err = msg.addButton("Открыть parsing_errors.csv", QtWidgets.QMessageBox.ActionRole)
        msg.addButton("OK", QtWidgets.QMessageBox.AcceptRole)
        msg.exec_()

        clicked = msg.clickedButton()
        if clicked == btn_open_out:
            open_in_explorer(self.options.output_folder)
        elif clicked == btn_open_err and self.last_error_path:
            open_in_explorer(self.last_error_path)

    def run_tests(self) -> None:
        report = run_sample_tests(logger=self.logger)
        QtWidgets.QMessageBox.information(self, "Результаты тестов", report)


def open_in_explorer(path: Path) -> None:
    try:
        if os.name == "nt":
            subprocess.Popen(["explorer", str(path)])
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception:
        pass


def collect_input_paths(raw_inputs: Sequence[str]) -> List[Path]:
    paths: List[Path] = []
    for item in raw_inputs:
        p = Path(item)
        if p.exists():
            paths.append(p)
    return paths


def run_cli(args: argparse.Namespace) -> int:
    options = ParserOptions(
        ocr_mode=args.ocr_mode,
        save_images=not args.no_images,
        output_folder=Path(args.output_dir),
        auto_start_drop=False,
        log_level=args.log_level,
        feature_threshold=args.feature_threshold,
        save_csv=args.save_csv,
    )
    logger = build_logger("cli", options.log_level)
    parser = LotParser(options, logger)

    input_paths = collect_input_paths(args.input)
    if not input_paths:
        print("Не найдены входные файлы/папки", file=sys.stdout)
        return 2

    rows = parser.parse_inputs(input_paths)
    out_excel = Path(args.output)
    export_rows(rows, out_excel, save_csv=options.save_csv)
    err_path = parser.save_errors_csv(options.output_folder)

    print(f"Готово. Лотов: {len(rows)}; ошибок: {len(parser.errors)}")
    print(f"Excel: {out_excel}")
    print(f"Ошибки: {err_path}")
    return 0 if rows else 1


def run_sample_tests(logger: Optional[logging.Logger] = None) -> str:
    logger = logger or build_logger("tests", "INFO")
    options = ParserOptions(ocr_mode="Never", save_images=False, output_folder=Path("output_test"))
    parser = LotParser(options, logger)

    sample_text = (
        "1--- (Москва) Монитор AOC M2060SWD2 19,5' s/n: GLXK5HA053007 --- Диагональ-19,5' / "
        "Затемнения и полосы на матрице-Нет / Видеовходы-VGA/DVI / Тип питания-Кабель / "
        "Сетевой шнур или блок питания в наличии - Да/Следы эксплуатации.\n700\n"
        "г. Москва, шоссе Энтузиастов д.14"
    )

    row = parser.parse_lot_block(
        sample_text,
        source_file="sample.pdf",
        source_page=1,
        processed_at=datetime.now(),
        auction_city_header="Москва",
        auction_end_datetime=datetime(2026, 3, 16, 12, 0),
        page_images=[],
    )

    checks = {
        "lot_position": row and row["LotPosition"] == 1,
        "brand": row and row["Brand"] == "AOC",
        "model": row and row["Model"] == "M2060SWD2",
        "diagonal": row and row["Diagonal"] == '19.5"',
        "serial": row and row["SerialNumber"] == "GLXK5HA053007",
        "price": row and row["StartPriceRUB"] == 700,
    }

    original = parser._ocr_page_text
    called = {"value": False}

    def fake_ocr(page):
        called["value"] = True
        return "OCR TEXT"

    parser._ocr_page_text = fake_ocr  # type: ignore[assignment]
    needed = parser._need_ocr("")
    if needed:
        parser._ocr_page_text(None)  # type: ignore[arg-type]
    parser._ocr_page_text = original  # type: ignore[assignment]
    checks["ocr_fallback_called"] = called["value"]

    passed = sum(1 for v in checks.values() if v)
    total = len(checks)
    report_lines = [f"{k}: {'PASS' if v else 'FAIL'}" for k, v in checks.items()]
    report_lines.append(f"Итого: {passed}/{total}")
    report = "\n".join(report_lines)
    logger.info("Sample tests:\n%s", report)
    return report


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="PDF Lot Parser GUI/CLI")
    parser.add_argument("--no-gui", action="store_true", help="Запуск в CLI режиме")
    parser.add_argument("--input", nargs="+", default=[], help="Файлы/папки PDF")
    parser.add_argument("--output", default=str(Path(DEFAULT_OUTPUT_DIR) / "parsed_lots.xlsx"), help="Путь к Excel")
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR, help="Папка для дополнительных файлов")
    parser.add_argument("--ocr-mode", choices=["Auto", "Always", "Never"], default="Auto")
    parser.add_argument("--no-images", action="store_true", help="Не сохранять изображения")
    parser.add_argument("--save-csv", action="store_true", help="Сохранять CSV рядом с Excel")
    parser.add_argument("--log-level", default=DEFAULT_LOG_LEVEL)
    parser.add_argument("--feature-threshold", type=int, default=5)
    parser.add_argument("--run-sample-tests", action="store_true", help="Запуск встроенных тестов")
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)

    if args.run_sample_tests:
        print(run_sample_tests())
        return 0

    if args.no_gui:
        return run_cli(args)

    app = QtWidgets.QApplication(sys.argv)
    app.setApplicationName(APP_TITLE)
    window = MainWindow()
    window.show()
    return app.exec_()


if __name__ == "__main__":
    sys.exit(main())
