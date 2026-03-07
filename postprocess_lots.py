#!/usr/bin/env python3
"""Пост-обработка выгрузки лотов из PDF в удобный CSV/XLSX для русскоязычного пользователя."""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path
from typing import Dict, Iterable, List

PHOTO_COLUMNS = [f"Photo{i}" for i in range(1, 14)]

RUS_COLUMN_MAP = {
    "SourceFile": "Файл_источник",
    "SourcePage": "Страница_источник",
    "ProcessedAt": "Время_обработки",
    "LotPosition": "Позиция_лота",
    "AuctionCity": "Город_торгов",
    "AuctionEndDateTime": "Окончание_торгов",
    "ItemType": "Тип_товара",
    "Brand": "Бренд",
    "Model": "Модель",
    "Diagonal": "Диагональ",
    "SerialNumber": "Серийный_номер",
    "StartPriceRUB": "Стартовая_цена_руб",
    "PickupAddress": "Адрес_самовывоза",
    "CityPickup": "Город_самовывоза",
    "VideoInputs": "Видеовходы",
    "PowerType": "Тип_питания",
    "PowerIncluded": "Питание_в_комплекте",
    "MatrixIssues": "Проблемы_матрицы",
    "Defect1": "Дефект_1",
    "Defect2": "Дефект_2",
    "HasImages": "Есть_изображения",
    "RawText": "Сырой_текст",
    "ExtractedFeatures": "Извлеченные_признаки",
    "Keywords": "Ключевые_слова",
    "EstimatedCondition": "Оценка_состояния",
    "Notes": "Заметки",
}

NOISE_PATTERNS = [
    re.compile(r"позиция\s*лота", re.IGNORECASE),
    re.compile(r"наименование\s+оборудования", re.IGNORECASE),
    re.compile(r"стартовая\s+цена", re.IGNORECASE),
    re.compile(r"город\s+и\s+место\s+получения", re.IGNORECASE),
    re.compile(r"фото\s*\d+", re.IGNORECASE),
]


def norm_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_lot_from_raw(raw: str) -> Dict[str, str]:
    txt = norm_text(raw)
    result: Dict[str, str] = {"Нормализованный_текст": txt}

    m = re.search(r"(?:^|\s)(\d{1,4})\s*---", txt)
    if m:
        result["Позиция_лота_из_текста"] = m.group(1)

    m = re.search(r"\(([^)]+)\)", txt)
    if m:
        result["Город_из_описания"] = m.group(1)

    m = re.search(r"\b(Монитор|Ноутбук|Системный блок|Телевизор|Принтер)\b", txt, re.IGNORECASE)
    if m:
        result["Тип_товара_из_текста"] = m.group(1).capitalize()

    m = re.search(r"\b([A-Za-z]{2,})\s+([A-Za-z0-9\-]{2,})\b", txt)
    if m:
        result["Бренд_из_текста"] = m.group(1).upper()
        result["Модель_из_текста"] = m.group(2)

    m = re.search(r"(?:s\/n|серийн(?:ый|ого)\s*номер)\s*[:\-]?\s*([A-Za-z0-9]{6,})", txt, re.IGNORECASE)
    if m:
        result["Серийный_номер_из_текста"] = m.group(1).upper()

    m = re.search(r"(\d{2}(?:[\.,]\d)?)\s*['\"]", txt)
    if m:
        result["Диагональ_из_текста"] = m.group(1).replace(",", ".")

    m = re.search(r"(\d{2,7})\s*(?:руб|₽)", txt, re.IGNORECASE)
    if m:
        result["Цена_из_текста_руб"] = m.group(1)

    return result


def is_noise_row(row: Dict[str, str]) -> bool:
    raw = norm_text(row.get("RawText", ""))
    lot = norm_text(row.get("LotPosition", ""))
    item = norm_text(row.get("ItemType", ""))
    serial = norm_text(row.get("SerialNumber", ""))
    price = norm_text(row.get("StartPriceRUB", ""))

    filled_main = sum(bool(v) for v in [lot, item, serial, price])

    if raw and sum(bool(p.search(raw)) for p in NOISE_PATTERNS) >= 2 and filled_main <= 1:
        return True
    if filled_main == 0 and len(raw) < 30:
        return True
    if lot == "" and item == "" and re.search(r"самовывоз|дата и время окончания", raw, re.IGNORECASE):
        return True
    return False


def fill_from_raw(row: Dict[str, str]) -> Dict[str, str]:
    parsed = parse_lot_from_raw(row.get("RawText", ""))
    merged = dict(row)
    merged.update(parsed)

    fill_pairs = [
        ("LotPosition", "Позиция_лота_из_текста"),
        ("AuctionCity", "Город_из_описания"),
        ("ItemType", "Тип_товара_из_текста"),
        ("Brand", "Бренд_из_текста"),
        ("Model", "Модель_из_текста"),
        ("SerialNumber", "Серийный_номер_из_текста"),
        ("Diagonal", "Диагональ_из_текста"),
        ("StartPriceRUB", "Цена_из_текста_руб"),
    ]
    for target, fallback in fill_pairs:
        if norm_text(merged.get(target, "")) == "" and norm_text(merged.get(fallback, "")) != "":
            merged[target] = merged[fallback]

    return merged


def score_row(row: Dict[str, str]) -> tuple[int, List[str]]:
    keys = [
        "LotPosition",
        "ItemType",
        "Brand",
        "Model",
        "SerialNumber",
        "StartPriceRUB",
        "PickupAddress",
        "AuctionEndDateTime",
    ]
    missing = [k for k in keys if norm_text(row.get(k, "")) == ""]
    score = round((len(keys) - len(missing)) / len(keys) * 100)
    return int(score), missing


def to_number(s: str) -> str:
    s = norm_text(s)
    return s if re.fullmatch(r"\d+(?:[\.,]\d+)?", s) else ""


def postprocess_rows(rows: List[Dict[str, str]]) -> List[Dict[str, str]]:
    # удаляем фото-колонки
    cleaned: List[Dict[str, str]] = []
    for row in rows:
        r = {k: ("" if v is None else str(v)) for k, v in row.items() if k not in PHOTO_COLUMNS}
        r = fill_from_raw(r)
        if not is_noise_row(r):
            score, missing = score_row(r)
            r["Полнота_заполнения_проц"] = str(score)
            r["Проблемные_поля"] = ", ".join(missing)
            cleaned.append(r)

    # дедупликация по файл+страница+позиция
    best_by_key: Dict[str, Dict[str, str]] = {}
    for row in cleaned:
        key = f"{norm_text(row.get('SourceFile'))}|{norm_text(row.get('SourcePage'))}|{norm_text(row.get('LotPosition'))}"
        current = best_by_key.get(key)
        if current is None or int(row.get("Полнота_заполнения_проц", "0")) > int(current.get("Полнота_заполнения_проц", "0")):
            best_by_key[key] = row

    result = list(best_by_key.values())

    # доп. полезные поля
    for row in result:
        row["Идентификатор_лота"] = (
            f"{norm_text(row.get('SourceFile'))}::{norm_text(row.get('SourcePage'))}::{norm_text(row.get('LotPosition'))}"
        )
        row["Наличие_серийного"] = "Да" if norm_text(row.get("SerialNumber")) else "Нет"
        row["Цена_числом_руб"] = to_number(row.get("StartPriceRUB", ""))

    # русификация названий
    rus_rows: List[Dict[str, str]] = []
    for row in result:
        renamed = {}
        for k, v in row.items():
            renamed[RUS_COLUMN_MAP.get(k, k)] = v
        # удаляем техполя парсинга
        for tmp in [
            "Позиция_лота_из_текста",
            "Город_из_описания",
            "Тип_товара_из_текста",
            "Бренд_из_текста",
            "Модель_из_текста",
            "Серийный_номер_из_текста",
            "Диагональ_из_текста",
            "Цена_из_текста_руб",
            "Нормализованный_текст",
        ]:
            renamed.pop(tmp, None)
        rus_rows.append(renamed)

    return rus_rows


def read_csv(path: Path) -> List[Dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: List[Dict[str, str]]) -> None:
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    all_keys = []
    seen = set()
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.add(key)
                all_keys.append(key)

    preferred = [
        "Идентификатор_лота",
        "Файл_источник",
        "Страница_источник",
        "Время_обработки",
        "Позиция_лота",
        "Город_торгов",
        "Окончание_торгов",
        "Тип_товара",
        "Бренд",
        "Модель",
        "Диагональ",
        "Серийный_номер",
        "Наличие_серийного",
        "Стартовая_цена_руб",
        "Цена_числом_руб",
        "Адрес_самовывоза",
        "Город_самовывоза",
        "Видеовходы",
        "Тип_питания",
        "Питание_в_комплекте",
        "Проблемы_матрицы",
        "Дефект_1",
        "Дефект_2",
        "Оценка_состояния",
        "Полнота_заполнения_проц",
        "Проблемные_поля",
        "Ключевые_слова",
        "Извлеченные_признаки",
        "Заметки",
        "Сырой_текст",
    ]
    ordered = [c for c in preferred if c in all_keys] + [c for c in all_keys if c not in preferred]

    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=ordered)
        writer.writeheader()
        writer.writerows(rows)


def maybe_read_xlsx(path: Path) -> List[Dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except Exception as exc:
        raise RuntimeError("Для XLSX установите openpyxl или используйте CSV") from exc

    wb = load_workbook(path, read_only=True)
    ws = wb.active
    data = list(ws.values)
    if not data:
        return []
    headers = ["" if h is None else str(h) for h in data[0]]
    rows: List[Dict[str, str]] = []
    for vals in data[1:]:
        row = {headers[i]: ("" if i >= len(vals) or vals[i] is None else str(vals[i])) for i in range(len(headers))}
        rows.append(row)
    return rows


def maybe_write_xlsx(path: Path, rows: List[Dict[str, str]]) -> None:
    try:
        from openpyxl import Workbook
    except Exception as exc:
        raise RuntimeError("Для XLSX установите openpyxl или используйте CSV") from exc

    wb = Workbook()
    ws = wb.active
    if not rows:
        wb.save(path)
        return

    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    wb.save(path)


def read_table(path: Path) -> List[Dict[str, str]]:
    if path.suffix.lower() == ".csv":
        return read_csv(path)
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        return maybe_read_xlsx(path)
    raise ValueError("Поддерживаются только CSV/XLSX")


def write_table(path: Path, rows: List[Dict[str, str]]) -> None:
    if path.suffix.lower() == ".csv":
        write_csv(path, rows)
    elif path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        maybe_write_xlsx(path, rows)
    else:
        raise ValueError("Поддерживаются только CSV/XLSX")


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="Пост-обработка выгрузки лотов")
    p.add_argument("input", type=Path)
    p.add_argument("output", type=Path)
    p.add_argument("--stats-json", type=Path, default=None)
    return p


def main() -> None:
    args = build_parser().parse_args()
    rows = read_table(args.input)
    out = postprocess_rows(rows)
    write_table(args.output, out)

    if args.stats_json:
        stats = {
            "rows_before": len(rows),
            "rows_after": len(out),
            "rows_removed": len(rows) - len(out),
            "columns_output": list(out[0].keys()) if out else [],
        }
        args.stats_json.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
